from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from typing import List, Optional
from .models import Project, Section
import os
from .config import settings

def generate_pptx(project: Project, sections: List[Section], image_paths: Optional[dict] = None) -> str:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Title slide with improved design
    title_slide_layout = prs.slide_layouts[6]  # Blank layout
    title_slide = prs.slides.add_slide(title_slide_layout)
    background = title_slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(25, 118, 210)  # Blue background
    
    # Add title
    title_box = title_slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    title_frame.text = project.title
    title_frame.paragraphs[0].font.size = Pt(54)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    title_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # Add subtitle
    subtitle_box = title_slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(1))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.word_wrap = True
    subtitle_frame.text = f"Topic: {project.topic}"
    subtitle_frame.paragraphs[0].font.size = Pt(24)
    subtitle_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    subtitle_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    image_paths = image_paths or {}
    
    for section in sorted(sections, key=lambda x: x.order):
        blank_slide_layout = prs.slide_layouts[6]  # Blank layout
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # Set background - white for cleaner look
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(255, 255, 255)
        
        # Add decorative top bar
        top_bar = slide.shapes.add_shape(
            1,  # Rectangle
            Inches(0), Inches(0), Inches(10), Inches(0.08)
        )
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = RGBColor(25, 118, 210)
        top_bar.line.color.rgb = RGBColor(25, 118, 210)
        
        # Add title with proper spacing
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.9))
        title_frame = title_box.text_frame
        title_frame.word_wrap = True
        title_frame.vertical_anchor = MSO_ANCHOR.TOP
        title_frame.text = section.title
        title_frame.paragraphs[0].font.size = Pt(32)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = RGBColor(25, 118, 210)
        
        # Determine if image will be added
        image_data = image_paths.get(section.id)
        has_image = image_data and os.path.exists(image_data.get("filepath", "") if isinstance(image_data, dict) else image_data)
        
        # Add content with proper spacing and alignment
        if section.content:
            if has_image:
                # Content on left, image on right - reduced height for better alignment
                content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(4.8), Inches(5.2))
            else:
                # Full width when no image - reduced height for better alignment
                content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(5.2))
            
            text_frame = content_box.text_frame
            text_frame.word_wrap = True
            text_frame.vertical_anchor = MSO_ANCHOR.TOP
            text_frame.margin_left = Inches(0.1)
            text_frame.margin_right = Inches(0.1)
            text_frame.margin_top = Inches(0.1)
            text_frame.margin_bottom = Inches(0.2)
            text_frame.text = section.content
            
            # Format paragraphs with better spacing
            for paragraph in text_frame.paragraphs:
                paragraph.font.size = Pt(14)
                paragraph.font.color.rgb = RGBColor(50, 50, 50)
                paragraph.space_before = Pt(4)
                paragraph.space_after = Pt(4)
                paragraph.line_spacing = 1.2  # Slightly tighter for better fit
                paragraph.alignment = PP_ALIGN.LEFT
        
        # Add image if available - positioned on the right with better alignment
        if has_image:
            try:
                # Place image on right side with proper spacing
                image_filepath = image_data.get("filepath", image_data) if isinstance(image_data, dict) else image_data
                pic = slide.shapes.add_picture(
                    image_filepath,
                    Inches(5.5),      # X position (right side, slightly more space)
                    Inches(1.4),      # Y position (below title with margin)
                    width=Inches(3.8)  # Slightly smaller width for better balance
                )
                # Ensure image doesn't extend beyond slide bounds
                if pic.height > Inches(5.0):
                    # Scale down if too tall
                    pic.height = Inches(5.0)
            except Exception as e:
                print(f"Failed to add image for section {section.id}: {str(e)}")
    
    filename = f"project_{project.id}_{project.title.replace(' ', '_')}.pptx"
    filepath = os.path.join(settings.EXPORT_DIR, filename)
    prs.save(filepath)
    
    return filepath
