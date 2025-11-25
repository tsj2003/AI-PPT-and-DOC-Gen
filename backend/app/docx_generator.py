from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from typing import List
from .models import Project, Section
import os
from .config import settings

def generate_docx(project: Project, sections: List[Section]) -> str:
    document = Document()
    
    title = document.add_heading(project.title, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    topic_para = document.add_paragraph()
    topic_run = topic_para.add_run(f"Topic: {project.topic}")
    topic_run.bold = True
    topic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    document.add_paragraph()
    
    for section in sorted(sections, key=lambda x: x.order):
        document.add_heading(section.title, level=1)
        
        if section.content:
            content_para = document.add_paragraph(section.content)
            content_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        document.add_paragraph()
    
    filename = f"project_{project.id}_{project.title.replace(' ', '_')}.docx"
    filepath = os.path.join(settings.EXPORT_DIR, filename)
    document.save(filepath)
    
    return filepath
